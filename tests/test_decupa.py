import json, os
from openpyxl import load_workbook
from comandos import decupa

PROBE = {"streams": [{"codec_type": "video", "width": 3840, "height": 2160, "r_frame_rate": "60000/1001",
                      "duration": "24.1", "codec_name": "h264",
                      "side_data_list": [{"rotation": -90}], "tags": {"timecode": "01:02:03:04"}},
                     {"codec_type": "audio", "channels": 2}],
         "format": {"duration": "24.1", "bit_rate": "100000000", "tags": {"creation_time": "2026-10-22T20:57:00Z"}}}
EXIF = {"XML:DeviceModelName": "ILCE-6700", "XML:AcquisitionRecordGroupItemValue": "s-log3-cine"}


def test_n_quadros_entre_6_e_18():
    assert [decupa.n_quadros(d) for d in (1, 24.1, 300)] == [6, 10, 18]


def test_montar_meta_gira_vertical_e_le_exif(tmp_path):
    f = tmp_path / "C4849.MP4"
    f.write_bytes(b"x")
    m = decupa.montar_meta(str(f), PROBE, EXIF)
    assert (m["clipe"], m["largura"], m["altura"], m["orientacao"]) == ("C4849", 2160, 3840, "vertical")
    assert (m["fps"], m["camera"], m["timecode"], m["audio_canais"]) == (59.94, "ILCE-6700", "01:02:03:04", 2)


def test_filtro_scale_antes_da_lut():
    assert decupa.filtro_video({"gamma": "s-log3-cine"}, "auto").index("scale") == 0
    assert "lut3d" in decupa.filtro_video({"gamma": "s-log3-cine"}, "auto")
    assert "lut3d" not in decupa.filtro_video({"gamma": None}, "auto")
    assert "lut3d" not in decupa.filtro_video({"gamma": "s-log3-cine"}, "nenhuma")
    assert "/x.cube" in decupa.filtro_video({}, "/x.cube")


def test_numeros_faltando():
    assert decupa.numeros_faltando(["C4849", "C4850", "C4853"]) == ["C4851", "C4852"]
    assert decupa.numeros_faltando(["DSC_0001", "DSC_0003"]) == ["DSC_0002"]
    assert decupa.numeros_faltando(["abc", "C1"]) == []


def test_fluxo_completo_sem_fala(pasta_clipes, capsys):
    p = str(pasta_clipes)
    dec = os.path.join(p, "_DECUPA")
    assert decupa.main(["preparar", p, "--sem-fala"]) == 0
    for n in ("c1", "c2", "c10"):
        assert os.path.exists(f"{dec}/meta/{n}.json")
        assert os.path.exists(f"{dec}/folhas/{n}.jpg")
        assert os.path.exists(f"{dec}/audio/{n}.wav")
        assert os.path.exists(f"{dec}/frames/{n}_06.jpg")          # 3 s = 6 quadros, pasta única
    assert "3 novos" in capsys.readouterr().out

    assert decupa.main(["preparar", p, "--sem-fala"]) == 0        # idempotente
    assert "3 já prontos" in capsys.readouterr().out

    with open(f"{dec}/descricoes/c2.json", "w") as f:
        json.dump({"clipe": "c2", "grupo": "Teste", "assunto": "barras", "aproveitamento": "Alto"}, f)
    assert decupa.main(["situacao", p, "--lote", "1"]) == 0
    assert json.load(open(f"{dec}/lotes.json")) == {"L1": ["c1"], "L2": ["c10"]}

    with open(f"{dec}/para_olhar.json", "w") as f:
        json.dump([{"assunto": "Marca", "clipes": "c2", "o_que_e": "logo"}], f)
    saida = os.path.join(p, "x.xlsx")
    assert decupa.main(["planilha", p, "--saida", saida, "--titulo", "TESTE"]) == 0
    wb = load_workbook(saida)
    assert wb.sheetnames == ["Decupagem", "Por grupo", "Para olhar", "Como ler"]
    ws = wb["Decupagem"]
    assert [ws.cell(row=r, column=2).value for r in (2, 3, 4)] == ["c1", "c2", "c10"]
    assert ws.cell(row=3, column=17).value == "Alto"
    assert len(ws._images) == 3
    assert "2 sem descrição" in capsys.readouterr().out


def test_situacao_antes_de_preparar(pasta_clipes, capsys):
    assert decupa.main(["situacao", str(pasta_clipes)]) == 0
    assert "faltam preparar 3" in capsys.readouterr().out


def test_planilha_sem_preparar_da_erro(pasta_clipes):
    assert decupa.main(["planilha", str(pasta_clipes)]) == 1


def test_nev_com_proxy_fica_de_fora(pasta_clipes):
    (pasta_clipes / "c1.NEV").write_bytes(b"x")
    assert [os.path.basename(c) for c in decupa.clipes_da_pasta(str(pasta_clipes))] == ["c1.mp4", "c2.mp4", "c10.mp4"]
